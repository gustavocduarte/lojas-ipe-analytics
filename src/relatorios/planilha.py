"""Montagem do .xlsx de fechamento (openpyxl): abas, formatos brasileiros, filtros e gráficos.

Os números entram como valores (nunca fórmulas), já calculados no dw: assim o arquivo é igual em
qualquer Excel e os testes conseguem lê-lo sem recalcular. Os formatos de número usam a notação
interna do Excel (vírgula de milhar, ponto decimal); quem abre em português do Brasil enxerga
"R$ 1.234,56" e "12,5%", e as datas aparecem como dd/mm/aaaa.
"""

from dataclasses import dataclass
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FORMATO_REAL = '"R$" #,##0.00'
FORMATO_PERCENTUAL = "0.0%"
FORMATO_DATA = "dd/mm/yyyy"
FORMATO_MES = "mm/yyyy"
FORMATO_INTEIRO = "#,##0"

COR_CABECALHO = "1F4E78"
COR_VERDE = "C6EFCE"
COR_VERMELHO = "FFC7CE"
COR_TEXTO_VERDE = "006100"
COR_TEXTO_VERMELHO = "9C0006"

ABA_RESUMO = "Resumo da rede"
ABA_LOJAS = "Lojas"
ABA_VENDEDORES = "Vendedores"
ABA_ALERTAS = "Alertas abertos"
ABA_PREVISAO = "Previsão"

# nome da coluna na consulta -> (título na planilha, formato)
COLUNAS_LOJAS = {
    "cod_loja": ("Loja", None), "nome_loja": ("Nome", None), "uf": ("UF", None),
    "venda_dia": ("Venda do dia", FORMATO_REAL), "meta_dia": ("Meta do dia", FORMATO_REAL),
    "cobertura_dia": ("Cobertura do dia", FORMATO_PERCENTUAL),
    "acumulado_mes": ("Acumulado do mês", FORMATO_REAL),
    "meta_acumulada": ("Meta acumulada até o dia", FORMATO_REAL),
    "meta_mes": ("Meta do mês", FORMATO_REAL),
    "cobertura_mes": ("Cobertura do mês", FORMATO_PERCENTUAL),
    "acumulado_ano_anterior": ("Mesmo período do ano anterior", FORMATO_REAL),
    "evolucao_ano_anterior": ("Evolução vs ano anterior", FORMATO_PERCENTUAL),
    "projecao_fechamento": ("Projeção de fechamento", FORMATO_REAL),
    "cobertura_projetada": ("Cobertura projetada", FORMATO_PERCENTUAL),
}  # fmt: skip
COLUNAS_VENDEDORES = {
    "posicao": ("Posição", FORMATO_INTEIRO), "cod_vendedor": ("Vendedor", None),
    "nome_vendedor": ("Nome", None), "cod_loja": ("Loja", None),
    "nome_loja": ("Nome da loja", None),
    "acumulado_mes": ("Acumulado do mês", FORMATO_REAL), "meta_mes": ("Meta do mês", FORMATO_REAL),
    "cobertura_mes": ("Cobertura do mês", FORMATO_PERCENTUAL),
}  # fmt: skip
COLUNAS_ALERTAS = {
    "severidade": ("Severidade", None), "tipo": ("Tipo", None), "data": ("Data", FORMATO_DATA),
    "loja_ou_centro": ("Loja ou centro de custo", None), "vendedor": ("Vendedor", None),
    "valor_envolvido": ("Valor envolvido", FORMATO_REAL), "status": ("Status", None),
    "explicacao": ("Explicação", None),
}  # fmt: skip
COLUNAS_PREVISAO = {
    "mes": ("Mês", FORMATO_MES), "cod_loja": ("Loja", None), "nome_loja": ("Nome", None),
    "previsto": ("Previsto", FORMATO_REAL), "faixa_baixa": ("Faixa baixa", FORMATO_REAL),
    "faixa_alta": ("Faixa alta", FORMATO_REAL),
    "wape_backtest_loja": ("Erro (WAPE) no backtest da loja", FORMATO_PERCENTUAL),
}  # fmt: skip


@dataclass
class DadosFechamento:
    """Tudo o que o relatório mostra, já consultado no dw."""

    mes: date  # primeiro dia do mês
    dia: date  # dia de posição
    lojas: pd.DataFrame
    indicadores: dict
    vendas_por_dia: pd.DataFrame
    vendedores: pd.DataFrame
    alertas: pd.DataFrame
    previsao_mensal: pd.DataFrame
    previsao_diaria: pd.DataFrame


def _valor(v):
    """NaN/NaT/None viram célula vazia; datas do pandas viram date."""
    if v is None or (not isinstance(v, str | date) and pd.isna(v)):
        return None
    return v.date() if isinstance(v, pd.Timestamp) else v


def _escrever_tabela(
    ws, quadro: pd.DataFrame, colunas: dict, linha_cabecalho: int, coluna: int = 1
):
    """Cabeçalho + linhas com formato por coluna; devolve a última linha escrita."""
    estilo = {
        "font": Font(bold=True, color="FFFFFF"),
        "fill": PatternFill("solid", fgColor=COR_CABECALHO),
        "alignment": Alignment(horizontal="center", vertical="center", wrap_text=True),
    }
    for j, (titulo, _) in enumerate(colunas.values()):
        celula = ws.cell(linha_cabecalho, coluna + j, titulo)
        celula.font, celula.fill, celula.alignment = (
            estilo["font"], estilo["fill"], estilo["alignment"],
        )  # fmt: skip
    ws.row_dimensions[linha_cabecalho].height = 32
    for i, linha in enumerate(quadro.itertuples(index=False), start=1):
        valores = dict(zip(quadro.columns, linha, strict=True))
        for j, (nome, (_, formato)) in enumerate(colunas.items()):
            celula = ws.cell(linha_cabecalho + i, coluna + j, _valor(valores[nome]))
            if formato:
                celula.number_format = formato
    for j, (nome, (titulo, _)) in enumerate(colunas.items()):
        maior = max([len(titulo) / 1.6] + [len(str(_valor(v) or "")) for v in quadro[nome]][:200])
        ws.column_dimensions[get_column_letter(coluna + j)].width = min(max(12, maior + 2), 60)
    return linha_cabecalho + len(quadro)


def _fixar_e_filtrar(ws, ultima_linha: int, ultima_coluna: int) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ultima_coluna)}{max(ultima_linha, 2)}"


def _verde_vermelho(ws, intervalo: str, limite: float = 1.0) -> None:
    """Verde a partir do limite (inclusive), vermelho abaixo dele."""
    verde = PatternFill("solid", bgColor=COR_VERDE, fgColor=COR_VERDE)
    vermelho = PatternFill("solid", bgColor=COR_VERMELHO, fgColor=COR_VERMELHO)
    ws.conditional_formatting.add(
        intervalo,
        CellIsRule(operator="greaterThanOrEqual", formula=[str(limite)], fill=verde,
                   font=Font(color=COR_TEXTO_VERDE)),
    )  # fmt: skip
    ws.conditional_formatting.add(
        intervalo,
        CellIsRule(operator="lessThan", formula=[str(limite)], fill=vermelho,
                   font=Font(color=COR_TEXTO_VERMELHO)),
    )  # fmt: skip


def _coluna_de(colunas: dict, nome: str) -> str:
    return get_column_letter(list(colunas).index(nome) + 1)


def _aba_lojas(wb: Workbook, dados: DadosFechamento) -> None:
    ws = wb.create_sheet(ABA_LOJAS)
    ultima = _escrever_tabela(ws, dados.lojas, COLUNAS_LOJAS, 1)
    _fixar_e_filtrar(ws, ultima, len(COLUNAS_LOJAS))
    for nome in ("cobertura_dia", "cobertura_mes", "cobertura_projetada"):
        letra = _coluna_de(COLUNAS_LOJAS, nome)
        _verde_vermelho(ws, f"{letra}2:{letra}{max(ultima, 2)}")
    letra = _coluna_de(COLUNAS_LOJAS, "evolucao_ano_anterior")
    _verde_vermelho(ws, f"{letra}2:{letra}{max(ultima, 2)}", 0.0)


def _aba_vendedores(wb: Workbook, dados: DadosFechamento) -> None:
    ws = wb.create_sheet(ABA_VENDEDORES)
    ultima = _escrever_tabela(ws, dados.vendedores, COLUNAS_VENDEDORES, 1)
    _fixar_e_filtrar(ws, ultima, len(COLUNAS_VENDEDORES))
    letra = _coluna_de(COLUNAS_VENDEDORES, "cobertura_mes")
    _verde_vermelho(ws, f"{letra}2:{letra}{max(ultima, 2)}")


def _aba_alertas(wb: Workbook, dados: DadosFechamento) -> None:
    ws = wb.create_sheet(ABA_ALERTAS)
    ultima = _escrever_tabela(ws, dados.alertas, COLUNAS_ALERTAS, 1)
    _fixar_e_filtrar(ws, ultima, len(COLUNAS_ALERTAS))
    ws.column_dimensions[_coluna_de(COLUNAS_ALERTAS, "explicacao")].width = 90
    cores = {"Crítica": "FF9999", "Alta": "FFC7CE", "Média": "FFEB9C", "Baixa": "E2EFDA"}
    coluna = _coluna_de(COLUNAS_ALERTAS, "severidade")
    for severidade, cor in cores.items():
        ws.conditional_formatting.add(
            f"{coluna}2:{coluna}{max(ultima, 2)}",
            CellIsRule(operator="equal", formula=[f'"{severidade}"'],
                       fill=PatternFill("solid", bgColor=cor, fgColor=cor)),
        )  # fmt: skip


def _aba_previsao(wb: Workbook, dados: DadosFechamento) -> None:
    ws = wb.create_sheet(ABA_PREVISAO)
    ultima = _escrever_tabela(ws, dados.previsao_mensal, COLUNAS_PREVISAO, 1)
    _fixar_e_filtrar(ws, ultima, len(COLUNAS_PREVISAO))
    diaria = dados.previsao_diaria
    if diaria.empty:
        return
    coluna = len(COLUNAS_PREVISAO) + 2
    titulos = {"dia": ("Dia", FORMATO_DATA), "previsto": ("Previsto da rede", FORMATO_REAL)}
    fim = _escrever_tabela(ws, diaria, titulos, 1, coluna)
    grafico = LineChart()
    grafico.title = "Previsão diária da rede"
    grafico.y_axis.title = "R$"
    grafico.y_axis.number_format = "#,##0"
    grafico.x_axis.number_format = "dd/mm"
    grafico.height, grafico.width = 8, 22
    grafico.add_data(
        Reference(ws, min_col=coluna + 1, min_row=1, max_row=fim), titles_from_data=True
    )
    grafico.set_categories(Reference(ws, min_col=coluna, min_row=2, max_row=fim))
    ws.add_chart(grafico, f"{get_column_letter(coluna + 3)}2")


def _aba_resumo(ws, dados: DadosFechamento) -> None:
    ws["A1"] = f"Lojas Ipê — fechamento de {dados.mes:%m/%Y}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        f"Posição em {dados.dia:%d/%m/%Y}. Valores em R$ (receita líquida de vendas concluídas)."
    )
    lojas, ind = dados.lojas, dados.indicadores
    soma = lambda coluna: float(lojas[coluna].fillna(0).sum())  # noqa: E731
    razao = lambda a, b: a / b if b else None  # noqa: E731
    venda_dia, meta_dia = soma("venda_dia"), soma("meta_dia")
    acumulado, meta_mes = soma("acumulado_mes"), soma("meta_mes")
    anterior = soma("acumulado_ano_anterior")
    evolucao = acumulado / anterior - 1 if anterior else None
    projecao = lojas["projecao_fechamento"].sum(min_count=1)
    projecao = None if pd.isna(projecao) else float(projecao)
    cobertura_projetada = None if projecao is None else razao(projecao, meta_mes)
    indicadores = [
        ("Venda do dia", venda_dia, FORMATO_REAL),
        ("Meta do dia", meta_dia, FORMATO_REAL),
        ("Cobertura do dia", razao(venda_dia, meta_dia), FORMATO_PERCENTUAL),
        ("Acumulado do mês", acumulado, FORMATO_REAL),
        ("Meta acumulada até o dia", soma("meta_acumulada"), FORMATO_REAL),
        ("Meta do mês", meta_mes, FORMATO_REAL),
        ("Cobertura do mês", razao(acumulado, meta_mes), FORMATO_PERCENTUAL),
        ("Mesmo período do ano anterior", anterior, FORMATO_REAL),
        ("Evolução vs ano anterior", evolucao, FORMATO_PERCENTUAL),
        ("Margem do mês", ind["margem"], FORMATO_REAL),
        ("Margem sobre a receita", ind["margem_percentual"], FORMATO_PERCENTUAL),
        ("Cupons do mês", ind["cupons"], FORMATO_INTEIRO),
        ("Ticket médio", ind["ticket_medio"], FORMATO_REAL),
        ("Projeção de fechamento", projecao, FORMATO_REAL),
        ("Cobertura projetada", cobertura_projetada, FORMATO_PERCENTUAL),
    ]  # fmt: skip
    for i, (rotulo, valor, formato) in enumerate(indicadores, start=4):
        ws.cell(i, 1, rotulo).font = Font(bold=True)
        celula = ws.cell(i, 2, valor)
        celula.number_format = formato
        if rotulo.startswith("Cobertura"):
            _verde_vermelho(ws, celula.coordinate)
        elif rotulo.startswith("Evolução"):
            _verde_vermelho(ws, celula.coordinate, 0.0)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 20

    primeira = 4 + len(indicadores) + 2
    ws.cell(primeira - 1, 1, "Venda por dia").font = Font(bold=True)
    titulos = {
        "dia": ("Dia", FORMATO_DATA), "venda": ("Venda", FORMATO_REAL),
        "meta_dia": ("Meta do dia", FORMATO_REAL),
    }  # fmt: skip
    fim = _escrever_tabela(ws, dados.vendas_por_dia, titulos, primeira)
    ws.column_dimensions["C"].width = 18
    barras = BarChart()
    barras.title = "Venda por dia x meta do dia"
    barras.y_axis.title = "R$"
    barras.y_axis.number_format = "#,##0"
    barras.x_axis.number_format = "dd"
    barras.height, barras.width = 9, 24
    barras.add_data(Reference(ws, min_col=2, min_row=primeira, max_row=fim), titles_from_data=True)
    categorias = Reference(ws, min_col=1, min_row=primeira + 1, max_row=fim)
    barras.set_categories(categorias)
    linha = LineChart()
    linha.add_data(Reference(ws, min_col=3, min_row=primeira, max_row=fim), titles_from_data=True)
    barras += linha
    ws.add_chart(barras, "E4")


def montar_pasta(dados: DadosFechamento, modelo: Workbook | None = None) -> Workbook:
    """Cria as cinco abas. Com `modelo` (ex.: o .xlsm com as macros), as abas antigas dele saem."""
    wb = modelo or Workbook()
    temporaria = wb.create_sheet("_temporaria", 0)
    for aba in [a for a in wb.worksheets if a is not temporaria]:
        wb.remove(aba)
    _aba_resumo(wb.create_sheet(ABA_RESUMO), dados)
    _aba_lojas(wb, dados)
    _aba_vendedores(wb, dados)
    _aba_alertas(wb, dados)
    _aba_previsao(wb, dados)
    wb.remove(temporaria)
    wb.active = 0
    return wb
