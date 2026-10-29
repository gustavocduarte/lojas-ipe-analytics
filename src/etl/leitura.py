"""Leitura das fontes (CSV ou Excel) conforme o mapeamento e carga no staging como texto cru.

Cada linha vira uma linha de stg.<entidade> com `arquivo_origem` (nome do arquivo),
`linha_origem` (1 = primeira linha de dados depois do cabeçalho, sem contar as ignoradas) e o
`id_carga`. Células vazias viram NULL. A carga em massa usa COPY, em lotes.
"""

import csv
import hashlib
import logging
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path

import psycopg
from openpyxl import load_workbook

from src import config
from src.etl import controle
from src.etl.contrato import CONTRATO, DATA, DECIMAL, HORA
from src.etl.mapeamento import Fonte, MapeamentoEntidade

log = logging.getLogger(__name__)

FORMATO_DATA_TEXTO = config.FORMATO_DATA_BR
BYTES_POR_BLOCO_HASH = 1 << 20


class ErroDeLeitura(RuntimeError):
    """O arquivo não pôde ser lido conforme o mapeamento (cabeçalho, aba ou coluna ausente)."""


@dataclass(frozen=True)
class ResultadoArquivo:
    entidade: str
    arquivo: str
    hash_arquivo: str
    linhas_lidas: int
    linhas_ignoradas: int
    situacao: str  # 'carregado' ou 'inalterado'


@dataclass(frozen=True)
class ResultadoStaging:
    id_carga: int
    arquivos: list[ResultadoArquivo]

    @property
    def entidades_carregadas(self) -> list[str]:
        return sorted({a.entidade for a in self.arquivos if a.situacao == "carregado"})

    @property
    def linhas_lidas(self) -> int:
        return sum(a.linhas_lidas for a in self.arquivos)


def hash_do_arquivo(caminho: Path) -> str:
    resumo = hashlib.sha256()
    with open(caminho, "rb") as arquivo:
        while bloco := arquivo.read(BYTES_POR_BLOCO_HASH):
            resumo.update(bloco)
    return resumo.hexdigest()


def arquivos_da_entidade(fonte: Fonte, mapa: MapeamentoEntidade) -> list[Path]:
    return sorted(p for p in fonte.pasta.glob(mapa.arquivo) if p.is_file())


# ---- conversão de células ----------------------------------------------------------------------
def celula_para_texto(valor, tipo: str, bruto: bool = False) -> str | None:
    """Texto de uma célula; vazio vira None e data/hora/número do Excel viram texto.

    `bruto=True` preserva os espaços das pontas (o staging é cru: quem limpa é o dw); sem ele o
    texto é aparado.
    """
    if valor is None:
        return None
    if isinstance(valor, datetime):
        if tipo == HORA:
            return valor.strftime(config.FORMATO_HORA)
        if tipo == DATA or valor.time() == time(0, 0):
            return valor.strftime(FORMATO_DATA_TEXTO)
        return valor.strftime(f"{FORMATO_DATA_TEXTO} {config.FORMATO_HORA}")
    if isinstance(valor, date):
        return valor.strftime(FORMATO_DATA_TEXTO)
    if isinstance(valor, time):
        return valor.strftime(config.FORMATO_HORA)
    if isinstance(valor, bool):
        return str(valor)
    if isinstance(valor, int):
        return str(valor)
    if isinstance(valor, float):
        return str(int(valor)) if valor.is_integer() else repr(valor)
    texto = str(valor)
    if not texto.strip():
        return None
    return texto if bruto else texto.strip()


def _ajustar_formatos(texto: str | None, tipo: str, mapa: MapeamentoEntidade) -> str | None:
    """Só mexe no texto quando o mapeamento declara um formato que não é o padrão do sistema."""
    if texto is None:
        return None
    if tipo == DATA and mapa.formato_data:
        try:
            return datetime.strptime(texto, mapa.formato_data).strftime(FORMATO_DATA_TEXTO)
        except ValueError:
            return texto  # fica cru; a limpeza decide (corrige ou rejeita com motivo)
    if tipo == DECIMAL and mapa.separador_decimal == ",":
        if mapa.separador_milhar:
            texto = texto.replace(mapa.separador_milhar, "")
        return texto.replace(",", ".")
    return texto


# ---- leitura das linhas ------------------------------------------------------------------------
def _linhas_csv(mapa: MapeamentoEntidade, caminho: Path) -> Iterator[list]:
    codificacao = "utf-8-sig" if mapa.codificacao.lower() in ("utf-8", "utf8") else mapa.codificacao
    with open(caminho, encoding=codificacao, newline="") as arquivo:
        yield from csv.reader(arquivo, delimiter=mapa.separador)


def _linhas_excel(mapa: MapeamentoEntidade, caminho: Path) -> Iterator[list]:
    pasta = load_workbook(caminho, read_only=True, data_only=True)
    try:
        if mapa.aba not in pasta.sheetnames:
            raise ErroDeLeitura(
                f"{caminho.name}: aba '{mapa.aba}' não encontrada (abas: {pasta.sheetnames})"
            )
        for linha in pasta[mapa.aba].iter_rows(values_only=True):
            yield list(linha)
    finally:
        pasta.close()


def ler_registros(
    mapa: MapeamentoEntidade, caminho: Path, contagem: dict | None = None
) -> Iterator[tuple]:
    """Linhas de dados como tuplas de texto (ou None), na ordem das colunas do mapeamento.

    `contagem` (opcional) recebe 'lidas' e 'ignoradas' ao final.
    """
    contagem = contagem if contagem is not None else {}
    contagem.update(lidas=0, ignoradas=0)
    contrato = CONTRATO[mapa.entidade]
    tipos = [contrato.coluna(destino).tipo for destino in mapa.colunas.values()]
    fonte_linhas = (
        _linhas_excel(mapa, caminho) if mapa.formato == "xlsx" else _linhas_csv(mapa, caminho)
    )

    indices: list[int] | None = None
    posicao_regras: list[tuple[int, object]] = []
    for numero, celulas in enumerate(fonte_linhas, start=1):
        if numero < mapa.linha_cabecalho:
            continue
        if numero == mapa.linha_cabecalho:
            cabecalho = [str(c).strip() if c is not None else "" for c in celulas]
            ausentes = [c for c in mapa.colunas if c not in cabecalho]
            if ausentes:
                raise ErroDeLeitura(
                    f"{caminho.name}: colunas do mapeamento ausentes no cabeçalho: {ausentes} "
                    f"(cabeçalho lido na linha {mapa.linha_cabecalho}: {cabecalho})"
                )
            indices = [cabecalho.index(c) for c in mapa.colunas]
            posicao_regras = [(cabecalho.index(r.coluna), r) for r in mapa.ignorar_linhas]
            continue

        if all(c is None or str(c).strip() == "" for c in celulas):
            contagem["ignoradas"] += 1
            continue
        if any(
            regra.casa(str(celulas[i]) if i < len(celulas) and celulas[i] is not None else "")
            for i, regra in posicao_regras
        ):
            contagem["ignoradas"] += 1
            continue
        contagem["lidas"] += 1
        yield tuple(
            _ajustar_formatos(
                celula_para_texto(celulas[i] if i < len(celulas) else None, tipo, bruto=True),
                tipo,
                mapa,
            )
            for i, tipo in zip(indices, tipos, strict=True)
        )
    if indices is None:
        raise ErroDeLeitura(
            f"{caminho.name}: arquivo sem a linha de cabeçalho ({mapa.linha_cabecalho})"
        )


# ---- carga no staging --------------------------------------------------------------------------
def _carregar_arquivo(
    conexao: psycopg.Connection,
    mapa: MapeamentoEntidade,
    caminho: Path,
    id_carga: int,
    tamanho_lote: int,
) -> tuple[int, int]:
    """Substitui no staging as linhas do arquivo (sem duplicar). Devolve (lidas, ignoradas)."""
    colunas = list(mapa.colunas.values())
    destino = ", ".join([*colunas, "arquivo_origem", "linha_origem", "id_carga"])
    conexao.execute(f"DELETE FROM stg.{mapa.entidade} WHERE arquivo_origem = %s", (caminho.name,))
    contagem: dict = {}
    registros = ler_registros(mapa, caminho, contagem)
    linha = 0
    with conexao.cursor() as cursor:
        while True:
            lote = []
            for registro in registros:
                linha += 1
                lote.append((*registro, caminho.name, linha, id_carga))
                if len(lote) >= tamanho_lote:
                    break
            if not lote:
                break
            with cursor.copy(f"COPY stg.{mapa.entidade} ({destino}) FROM STDIN") as copia:
                for registro in lote:
                    copia.write_row(registro)
    return contagem["lidas"], contagem["ignoradas"]


def carregar_para_staging(
    conexao: psycopg.Connection,
    fonte: Fonte,
    *,
    entidades: list[str] | None = None,
    modo: str = "incremental",
    tamanho_lote: int = config.TAMANHO_LOTE_CARGA,
) -> ResultadoStaging:
    """Lê os arquivos da fonte para o staging dentro de uma nova carga.

    No modo incremental, arquivos com o mesmo conteúdo (hash) de uma carga já concluída são
    pulados; no modo completo, todos são relidos. O staging é confirmado (commit) ao final;
    arquivos só passam a contar como processados depois do dw (ver controle.confirmar_arquivos).
    """
    escolhidas = entidades or list(fonte.entidades)
    desconhecidas = set(escolhidas) - set(fonte.entidades)
    if desconhecidas:
        raise ErroDeLeitura(
            f"entidades não mapeadas na fonte '{fonte.nome}': {sorted(desconhecidas)}"
        )

    id_carga = controle.iniciar_carga(conexao, fonte.nome, modo)
    resultados: list[ResultadoArquivo] = []
    try:
        for entidade in escolhidas:
            mapa = fonte.entidades[entidade]
            arquivos = arquivos_da_entidade(fonte, mapa)
            if not arquivos:
                log.warning(
                    "fonte %s, %s: nenhum arquivo casa com '%s'", fonte.nome, entidade, mapa.arquivo
                )
            for caminho in arquivos:
                resumo = hash_do_arquivo(caminho)
                if modo == "incremental" and controle.arquivo_ja_processado(
                    conexao, fonte.nome, entidade, caminho.name, resumo
                ):
                    resultado = ResultadoArquivo(entidade, caminho.name, resumo, 0, 0, "inalterado")
                    log.info("%s: %s inalterado, não relido", entidade, caminho.name)
                else:
                    lidas, ignoradas = _carregar_arquivo(
                        conexao, mapa, caminho, id_carga, tamanho_lote
                    )
                    resultado = ResultadoArquivo(
                        entidade, caminho.name, resumo, lidas, ignoradas, "carregado"
                    )
                    log.info("%s: %s -> %d linhas no staging", entidade, caminho.name, lidas)
                controle.registrar_arquivo(conexao, id_carga, resultado)
                resultados.append(resultado)
        conexao.commit()
    except Exception as erro:
        conexao.rollback()
        controle.finalizar_carga(conexao, id_carga, "erro", str(erro))
        raise
    return ResultadoStaging(id_carga, resultados)
