"""Segundo exemplo de fonte: uma planilha Excel de "outra empresa" com as vendas de um mês.

Serve para provar que o ETL é genérico: os nomes de coluna são outros ("Nº Cupom", "Dt. Emissão",
"Vlr. Líquido"), o cabeçalho está na linha 3, há uma linha de total no fim e uma aba extra que deve
ser ignorada. Os valores são os mesmos de um mês do ERP, em células nativas do Excel (datas, horas
e números), para que o mapeamento próprio (config/fontes/outra_empresa.yaml) carregue exatamente o
mesmo conteúdo.
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from src import config

ABA_VENDAS = "Vendas"
ABA_EXTRA = "Resumo"
LINHA_CABECALHO = 3
ROTULO_TOTAL = "TOTAL"
EMPRESA = "Rede Exemplo Comercial Ltda."

# (coluna do sistema, cabeçalho da outra empresa)
COLUNAS = (
    ("numero_cupom", "Nº Cupom"),
    ("numero_item", "Item"),
    ("id_venda_erp", "Cód. Lançamento"),
    ("data_venda", "Dt. Emissão"),
    ("hora_venda", "Hora"),
    ("cod_loja", "Filial"),
    ("caixa", "PDV"),
    ("cod_vendedor", "Vendedor"),
    ("cod_cliente", "Cliente"),
    ("cod_forma_pagamento", "Forma Pgto"),
    ("numero_parcelas", "Parcelas"),
    ("cod_campanha", "Campanha"),
    ("sku", "Produto (SKU)"),
    ("quantidade", "Qtde"),
    ("preco_unitario", "Preço Unit."),
    ("desconto", "Desc. R$"),
    ("receita_liquida", "Vlr. Líquido"),
    ("custo_total", "Custo Total"),
    ("total_cupom", "Total do Cupom"),
    ("status", "Situação"),
    ("motivo_cancelamento_devolucao", "Motivo"),
)
COLUNAS_TEXTO_VAZIO_VIRA_NULO = ("cod_cliente", "cod_campanha", "motivo_cancelamento_devolucao")
COLUNAS_INTEIRAS = ("numero_item", "caixa", "numero_parcelas", "quantidade")
COLUNAS_DECIMAIS = (
    "preco_unitario",
    "desconto",
    "receita_liquida",
    "custo_total",
    "total_cupom",
)


def _valor_da_celula(coluna: str, valor):
    if coluna == "data_venda":
        return valor.date()
    if coluna == "hora_venda":
        return valor.time()
    if coluna in COLUNAS_INTEIRAS:
        return int(valor)
    if coluna in COLUNAS_DECIMAIS:
        return float(valor)
    if coluna in COLUNAS_TEXTO_VAZIO_VIRA_NULO and valor == "":
        return None
    return str(valor)


def gravar_vendas(df: pd.DataFrame, caminho: Path, ano: int, mes: int) -> int:
    """Grava as vendas do mês (DataFrame tipado do gerador) no layout da outra empresa.

    Devolve o número de linhas de dados (sem cabeçalho nem total).
    """
    caminho.parent.mkdir(parents=True, exist_ok=True)
    pasta = Workbook(write_only=True)
    resumo = pasta.create_sheet(ABA_EXTRA)  # vem primeiro de propósito: a aba certa é a segunda
    resumo.append(["Esta aba não faz parte da carga e deve ser ignorada."])
    resumo.append(["Linhas de venda no mês", len(df)])

    aba = pasta.create_sheet(ABA_VENDAS)
    aba.append([f"{EMPRESA} - Vendas de {config.NOMES_MESES[mes - 1].lower()}/{ano}"])
    aba.append([])
    aba.append([cabecalho for _, cabecalho in COLUNAS])
    nomes = [coluna for coluna, _ in COLUNAS]
    for linha in df[nomes].itertuples(index=False, name=None):
        aba.append([_valor_da_celula(c, v) for c, v in zip(nomes, linha, strict=True)])
    aba.append([])
    total = [None] * len(COLUNAS)
    total[0] = ROTULO_TOTAL
    total[nomes.index("receita_liquida")] = float(round(df["receita_liquida"].sum(), 2))
    aba.append(total)
    pasta.save(caminho)
    return len(df)
