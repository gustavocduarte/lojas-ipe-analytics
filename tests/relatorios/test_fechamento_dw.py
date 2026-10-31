"""O .xlsx de fechamento precisa bater com consultas independentes ao dw (precisa de PostgreSQL)."""

from datetime import date, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from src.relatorios import fechamento
from src.relatorios.__main__ import principal
from src.relatorios.planilha import ABA_ALERTAS, ABA_LOJAS, ABA_PREVISAO, ABA_RESUMO, ABA_VENDEDORES
from tests.etl import base_alertas as b

pytestmark = pytest.mark.banco

MES, DIA = date(2025, 3, 1), date(2025, 3, 15)


def _dim_tempo(con, inicio: date, fim: date):
    dia = inicio
    while dia <= fim:
        con.execute(
            "INSERT INTO dw.dim_tempo (sk_tempo, data, dia, semana, mes, trimestre, ano, nome_mes,"
            " nome_dia_semana, feriado_nacional)"
            " VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, FALSE)",
            (
                int(dia.strftime("%Y%m%d")), dia, dia.day, dia.isocalendar()[1], dia.month,
                (dia.month - 1) // 3 + 1, dia.year, b.MESES[dia.month - 1], b.DIAS[dia.weekday()],
            ),
        )  # fmt: skip
        dia += timedelta(days=1)


@pytest.fixture
def dw(banco_com_ddl):
    con = banco_com_ddl
    b.montar_base(con, ano=2025)
    _dim_tempo(con, date(2024, 3, 1), date(2024, 3, 31))
    n = 0
    for dia in range(1, 21):  # março de 2025 até o dia 20 (depois do dia de posição, de propósito)
        for loja, base in (("L01", 100), ("L02", 50)):
            n += 1
            b.item(con, f"A{n}", cupom=f"A{n}", data=f"2025-03-{dia:02d}", loja=loja,
                   vendedor="V1" if loja == "L01" else "V2", receita=base * dia)  # fmt: skip
    for dia in range(1, 16):  # março de 2024, mesmo período
        b.item(con, f"P{dia}", cupom=f"P{dia}", data=f"2024-03-{dia:02d}", receita=80 * dia)
    b.item(con, "CANC", cupom="CANC", data="2025-03-05", receita=9999, status="Cancelada")
    b.item(con, "OUTRO", cupom="OUTRO", data="2025-04-02", receita=7777)  # outro mês
    for cod, valor in (("L01", 6000), ("L02", 4000)):
        con.execute(
            "INSERT INTO dw.fato_meta_vendas (sk_tempo, sk_loja, setor, valor_meta)"
            " SELECT 20250301, sk_loja, 'Alimentos', %s FROM dw.dim_loja WHERE cod_loja = %s",
            (valor, cod),
        )
    for cod, valor in (("V1", 5000), ("V2", 3000)):
        con.execute(
            "INSERT INTO dw.fato_meta_vendedor (sk_tempo, sk_vendedor, valor_meta)"
            " SELECT 20250301, sk_vendedor, %s FROM dw.dim_vendedor WHERE cod_vendedor = %s",
            (valor, cod),
        )
    # alertas: desconto muito acima do limite (um novo, um já classificado como falso positivo)
    b.item(con, "DESC1", cupom="DESC1", data="2025-03-07", preco=100, desconto=40, receita=60)
    b.item(con, "DESC2", cupom="DESC2", data="2025-03-08", preco=100, desconto=45, receita=55)
    con.commit()
    b.gerar(con)
    con.execute(
        "UPDATE dw.alertas SET status = 'falso_positivo' WHERE sk_alerta ="
        " (SELECT min(sk_alerta) FROM dw.alertas WHERE tipo_alerta = 'desconto_acima_limite')"
    )
    # previsão (como se tivesse sido gerada em 31/12/2025)
    con.execute(
        "INSERT INTO dw.previsao_vendas (granularidade, data, sk_loja, setor, valor_previsto,"
        " valor_baixo, valor_alto, versao_modelo, treino_ate)"
        " SELECT 'mes', '2026-01-01', sk_loja, 'Alimentos', 20000, 17000, 23000, 'x@1',"
        " '2025-12-31'"
        " FROM dw.dim_loja"
    )
    con.execute(
        "INSERT INTO dw.previsao_vendas (granularidade, data, sk_loja, setor, valor_previsto,"
        " valor_baixo, valor_alto, versao_modelo, treino_ate)"
        " SELECT 'dia', make_date(2026, 1, d), sk_loja, 'Alimentos', 600, 500, 700, 'x@1',"
        " '2025-12-31' FROM dw.dim_loja, generate_series(1, 31) d"
    )
    con.execute(
        "INSERT INTO dw.previsao_backtest (versao_modelo, nivel, granularidade, sk_loja,"
        " treino_ate,"
        " teste_inicio, teste_fim, wape, vies, escolhido)"
        " SELECT 'x@1', 'loja', 'mes', sk_loja, '2025-06-30', '2025-07-01', '2025-12-31', 0.1234,"
        " 0, TRUE FROM dw.dim_loja WHERE cod_loja = 'L01'"
    )
    con.commit()
    return con


@pytest.fixture
def pasta(dw, tmp_path):
    caminho, _ = fechamento.gerar(dw, MES, DIA, tmp_path / "f.xlsx", modelo=None)
    return load_workbook(caminho)


def _escalar(con, sql, *params):
    valor = con.execute(sql, params).fetchone()[0]
    return float(valor) if isinstance(valor, Decimal) else valor


def _tabela(ws):
    titulos = [c.value for c in ws[1]]
    return [
        dict(zip(titulos, [c.value for c in linha], strict=True))
        for linha in ws.iter_rows(min_row=2)
    ]


class TestLojas:
    def test_venda_acumulado_e_meta_de_cada_loja_batem_com_o_dw(self, dw, pasta):
        linhas = {linha["Loja"]: linha for linha in _tabela(pasta[ABA_LOJAS])}
        assert set(linhas) == {"L01", "L02"}
        for cod in ("L01", "L02"):
            sql = (
                "SELECT {} FROM dw.fato_vendas v JOIN dw.dim_tempo t USING (sk_tempo)"
                " JOIN dw.dim_loja l USING (sk_loja) WHERE l.cod_loja = %s"
                " AND v.status = 'Concluída'"
                " AND t.data {}"
            )
            acumulado = _escalar(
                dw,
                sql.format("sum(v.receita_liquida)", "BETWEEN '2025-03-01' AND '2025-03-15'"),
                cod,
            )
            dia = _escalar(dw, sql.format("sum(v.receita_liquida)", "= '2025-03-15'"), cod)
            anterior = _escalar(
                dw,
                sql.format("sum(v.receita_liquida)", "BETWEEN '2024-03-01' AND '2024-03-15'"),
                cod,
            )
            meta = _escalar(
                dw, "SELECT sum(valor_meta) FROM dw.fato_meta_vendas m JOIN dw.dim_loja l"
                " USING (sk_loja) WHERE l.cod_loja = %s AND m.sk_tempo = 20250301", cod,
            )  # fmt: skip
            linha = linhas[cod]
            assert linha["Acumulado do mês"] == pytest.approx(acumulado)
            assert linha["Venda do dia"] == pytest.approx(dia)
            assert linha["Meta do mês"] == pytest.approx(meta)
            assert linha["Cobertura do mês"] == pytest.approx(acumulado / meta)
            assert linha["Mesmo período do ano anterior"] == (
                pytest.approx(anterior) if anterior else None
            )
            if anterior:
                assert linha["Evolução vs ano anterior"] == pytest.approx(acumulado / anterior - 1)

    def test_valores_conhecidos_a_mao(self, pasta):
        l01 = next(x for x in _tabela(pasta[ABA_LOJAS]) if x["Loja"] == "L01")
        assert l01["Acumulado do mês"] == pytest.approx(100 * sum(range(1, 16)) + 60 + 55)
        assert l01["Mesmo período do ano anterior"] == pytest.approx(80 * sum(range(1, 16)))
        assert l01["Venda do dia"] == 1500

    def test_meta_do_dia_vem_da_meta_diaria_e_soma_com_a_do_mes(self, dw, pasta):
        soma_ate_dia = _escalar(
            dw, "SELECT sum(valor_meta_dia) FROM dw.vw_meta_vendas_diaria"
            " WHERE data BETWEEN '2025-03-01' AND '2025-03-15'",
        )  # fmt: skip
        linhas = _tabela(pasta[ABA_LOJAS])
        assert sum(x["Meta acumulada até o dia"] for x in linhas) == pytest.approx(soma_ate_dia)
        assert all(x["Meta do dia"] > 0 for x in linhas)


class TestResumo:
    def test_totais_da_rede_batem_com_o_dw(self, dw, pasta):
        ws = pasta[ABA_RESUMO]
        valores = {ws.cell(i, 1).value: ws.cell(i, 2).value for i in range(4, 20)}
        receita = _escalar(
            dw, "SELECT sum(receita_liquida) FROM dw.fato_vendas v JOIN dw.dim_tempo t USING"
            " (sk_tempo) WHERE status = 'Concluída'"
            " AND t.data BETWEEN '2025-03-01' AND '2025-03-15'",
        )  # fmt: skip
        margem = _escalar(
            dw, "SELECT sum(margem) FROM dw.fato_vendas v JOIN dw.dim_tempo t USING (sk_tempo)"
            " WHERE status = 'Concluída' AND t.data BETWEEN '2025-03-01' AND '2025-03-15'",
        )  # fmt: skip
        cupons = _escalar(
            dw, "SELECT count(DISTINCT (sk_loja, numero_cupom)) FROM dw.fato_vendas v JOIN"
            " dw.dim_tempo t USING (sk_tempo) WHERE status = 'Concluída'"
            " AND t.data BETWEEN '2025-03-01' AND '2025-03-15'",
        )  # fmt: skip
        meta = _escalar(
            dw, "SELECT sum(valor_meta) FROM dw.fato_meta_vendas WHERE sk_tempo = 20250301"
        )
        assert valores["Acumulado do mês"] == pytest.approx(receita)
        assert valores["Margem do mês"] == pytest.approx(margem)
        assert valores["Cupons do mês"] == cupons
        assert valores["Ticket médio"] == pytest.approx(receita / cupons)
        assert valores["Meta do mês"] == pytest.approx(meta)
        assert valores["Cobertura do mês"] == pytest.approx(receita / meta)

    def test_venda_por_dia_ignora_canceladas_e_para_no_dia_de_posicao(self, pasta):
        ws = pasta[ABA_RESUMO]
        inicio = next(i for i in range(1, 60) if ws.cell(i, 1).value == "Dia")
        dias = [(ws.cell(inicio + k, 2).value, ws.cell(inicio + k, 3).value) for k in range(1, 32)]
        assert dias[0][0] == pytest.approx(150)  # 100 + 50 no dia 1
        assert dias[14][0] == pytest.approx(150 * 15)
        assert dias[4][0] == pytest.approx(150 * 5)  # a cancelada de 9.999 não entra
        assert all(venda is None for venda, _ in dias[15:])  # depois do dia 15 não há venda
        assert all(meta and meta > 0 for _, meta in dias)


class TestVendedoresAlertasPrevisao:
    def test_ranking_de_vendedores(self, dw, pasta):
        linhas = _tabela(pasta[ABA_VENDEDORES])
        assert [x["Vendedor"] for x in linhas] == ["V1", "V2"]
        assert [x["Posição"] for x in linhas] == [1, 2]
        for x in linhas:
            esperado = _escalar(
                dw, "SELECT sum(receita_liquida) FROM dw.fato_vendas v JOIN dw.dim_tempo t USING"
                " (sk_tempo) JOIN dw.dim_vendedor d USING (sk_vendedor) WHERE d.cod_vendedor = %s"
                " AND status = 'Concluída' AND t.data BETWEEN '2025-03-01' AND '2025-03-15'",
                x["Vendedor"],
            )  # fmt: skip
            assert x["Acumulado do mês"] == pytest.approx(esperado)
        assert linhas[0]["Meta do mês"] == 5000

    def test_so_alertas_independentes_abertos_e_do_mais_grave_ao_menos(self, dw, pasta):
        esperado = _escalar(
            dw, "SELECT count(*) FROM dw.alertas WHERE vigente AND independente"
            " AND status IN ('novo', 'em_analise')",
        )  # fmt: skip
        linhas = _tabela(pasta[ABA_ALERTAS])
        assert len(linhas) == esperado >= 1
        assert all(x["Status"] in ("Novo", "Em análise") for x in linhas)
        ordem = {"Crítica": 4, "Alta": 3, "Média": 2, "Baixa": 1}
        graus = [ordem[x["Severidade"]] for x in linhas]
        assert graus == sorted(graus, reverse=True)
        falsos = _escalar(dw, "SELECT count(*) FROM dw.alertas WHERE status = 'falso_positivo'")
        assert falsos == 1  # existe um falso positivo e ele ficou de fora

    def test_previsao_mensal_e_diaria_da_execucao_mais_recente(self, dw, pasta):
        ws = pasta[ABA_PREVISAO]
        linhas = _tabela(ws)
        assert len([x for x in linhas if x["Loja"]]) == 2
        l01 = next(x for x in linhas if x["Loja"] == "L01")
        assert (
            l01["Previsto"] == 20000 and l01["Faixa baixa"] == 17000 and l01["Faixa alta"] == 23000
        )
        assert l01["Erro (WAPE) no backtest da loja"] == pytest.approx(0.1234)
        l02 = next(x for x in linhas if x["Loja"] == "L02")
        assert l02["Erro (WAPE) no backtest da loja"] is None
        previsto_dia = [ws.cell(i, 10).value for i in range(2, 33)]
        assert previsto_dia == [1200.0] * 31  # 2 lojas x 600
        assert len(ws._charts) == 1


class TestComando:
    def test_gera_o_arquivo_padrao_e_informa_o_caminho(self, dw, tmp_path, capsys):
        saida = tmp_path / "sub" / "fechamento.xlsx"
        codigo = principal(
            ["fechamento", "--mes", "2025-03", "--dia", "2025-03-15", "--saida", str(saida)],
            conexao=dw, verificar_memoria=None,
        )  # fmt: skip
        assert codigo == 0 and saida.exists()
        texto = capsys.readouterr().out
        assert "fechamento de 03/2025 (posição 15/03/2025)" in texto and "sem macros" in texto

    def test_sem_argumentos_usa_o_ultimo_dia_com_venda(self, dw, tmp_path, capsys):
        saida = tmp_path / "f.xlsx"
        assert (
            principal(["fechamento", "--saida", str(saida)], conexao=dw, verificar_memoria=None)
            == 0
        )
        assert "posição 02/04/2025" in capsys.readouterr().out

    def test_dia_fora_do_mes_sai_com_2(self, dw, tmp_path):
        argumentos = ["fechamento", "--mes", "2025-03", "--dia", "2025-04-02"]
        assert principal(argumentos, conexao=dw, verificar_memoria=None) == 2

    def test_extensao_xlsm_sem_modelo_sai_com_2(self, dw, tmp_path):
        argumentos = ["fechamento", "--sem-macros", "--saida", str(tmp_path / "f.xlsm")]
        assert principal(argumentos, conexao=dw, verificar_memoria=None) == 2

    def test_data_invalida_sai_com_2(self, dw):
        assert principal(["fechamento", "--dia", "ontem"], conexao=dw, verificar_memoria=None) == 2


class TestModeloDeMacros:
    """Com um modelo .xlsm presente a saída é .xlsm; o modelo de verdade é testado em tests/vba."""

    @pytest.fixture
    def modelo(self, tmp_path):
        from openpyxl import Workbook

        caminho = tmp_path / "modelo.xlsm"
        Workbook().save(caminho)  # sem VBA de verdade: só exercita o caminho de código
        return caminho

    def test_com_modelo_a_saida_e_xlsm_e_traz_as_cinco_abas(self, dw, tmp_path, modelo):
        destino, desempenho = fechamento.gerar(dw, MES, DIA, tmp_path / "f.xlsm", modelo=modelo)
        assert desempenho["com_macros"] is True
        assert load_workbook(destino, keep_vba=True).sheetnames[0] == ABA_RESUMO

    def test_caminho_padrao_depende_do_modelo(self, tmp_path):
        assert fechamento.caminho_padrao(MES, False, tmp_path).name == "fechamento_2025-03.xlsx"
        assert fechamento.caminho_padrao(MES, True, tmp_path).name == "fechamento_2025-03.xlsm"

    def test_com_modelo_saida_xlsx_e_recusada(self, dw, tmp_path, modelo):
        with pytest.raises(ValueError, match="xlsm"):
            fechamento.gerar(dw, MES, DIA, tmp_path / "f.xlsx", modelo=modelo)

    def test_modelo_inexistente_gera_xlsx_sem_macros(self, dw, tmp_path):
        destino, desempenho = fechamento.gerar(
            dw, MES, DIA, tmp_path / "f.xlsx", modelo=tmp_path / "nao_existe.xlsm"
        )
        assert desempenho["com_macros"] is False and destino.suffix == ".xlsx"
