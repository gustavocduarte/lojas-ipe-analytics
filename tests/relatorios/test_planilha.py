"""Estrutura e formatação do .xlsx de fechamento, com dados sintéticos (sem banco)."""

from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.relatorios import planilha
from src.relatorios.planilha import DadosFechamento, montar_pasta


def _dados(com_previsao: bool = True, com_projecao: bool = True) -> DadosFechamento:
    lojas = pd.DataFrame(
        {
            "cod_loja": ["L01", "L02"], "nome_loja": ["Ipê 1", "Ipê 2"], "uf": ["SP", "RJ"],
            "venda_dia": [1500.0, 400.0], "meta_dia": [1000.0, 800.0],
            "cobertura_dia": [1.5, 0.5], "acumulado_mes": [12000.0, 6000.0],
            "meta_acumulada": [9000.0, 8000.0], "meta_mes": [20000.0, 16000.0],
            "cobertura_mes": [0.6, 0.375], "acumulado_ano_anterior": [10000.0, 7000.0],
            "evolucao_ano_anterior": [0.2, -0.142857],
            "projecao_fechamento": [24000.0, 13000.0] if com_projecao else [None, None],
            "cobertura_projetada": [1.2, 0.8125] if com_projecao else [None, None],
        }
    )  # fmt: skip
    dias = pd.date_range("2025-03-01", "2025-03-31")
    vendas = pd.DataFrame(
        {
            "dia": dias.date,
            "venda": [100.0 * i if i <= 15 else None for i in range(1, 32)],
            "meta_dia": [1000.0] * 31,
        }
    )
    vendedores = pd.DataFrame(
        {
            "posicao": [1, 2], "cod_vendedor": ["V1", "V2"], "nome_vendedor": ["Ana", "Bia"],
            "cod_loja": ["L01", "L02"], "nome_loja": ["Ipê 1", "Ipê 2"],
            "acumulado_mes": [12000.0, 6000.0], "meta_mes": [10000.0, 9000.0],
            "cobertura_mes": [1.2, 0.6667],
        }
    )  # fmt: skip
    alertas = pd.DataFrame(
        {
            "severidade": ["Crítica", "Baixa"], "tipo": ["Desconto acima do limite"] * 2,
            "data": [date(2025, 3, 5), date(2025, 3, 6)], "loja_ou_centro": ["Ipê 1", "Ipê 2"],
            "vendedor": ["Ana", None], "valor_envolvido": [350.5, 20.0],
            "status": ["Novo", "Em análise"], "explicacao": ["Desconto de 40%", "Desconto de 19%"],
        }
    )  # fmt: skip
    previsao = pd.DataFrame(
        {
            "mes": [date(2026, 1, 1), date(2026, 1, 1)], "cod_loja": ["L01", "L02"],
            "nome_loja": ["Ipê 1", "Ipê 2"], "previsto": [20000.0, 15000.0],
            "faixa_baixa": [17000.0, 12000.0], "faixa_alta": [23000.0, 18000.0],
            "wape_backtest_loja": [0.11, 0.14], "versao_modelo": ["x@1"] * 2,
            "treino_ate": [date(2025, 12, 31)] * 2,
        }
    )  # fmt: skip
    diaria = pd.DataFrame(
        {"dia": pd.date_range("2026-01-01", periods=31).date, "previsto": [1000.0] * 31}
    )
    return DadosFechamento(
        mes=date(2025, 3, 1), dia=date(2025, 3, 15), lojas=lojas,
        indicadores={"receita": 18000.0, "margem": 7200.0, "margem_percentual": 0.4,
                     "cupons": 60, "ticket_medio": 300.0},
        vendas_por_dia=vendas, vendedores=vendedores, alertas=alertas,
        previsao_mensal=previsao if com_previsao else previsao.iloc[0:0],
        previsao_diaria=diaria if com_previsao else diaria.iloc[0:0],
    )  # fmt: skip


@pytest.fixture
def pasta(tmp_path):
    caminho = tmp_path / "f.xlsx"
    montar_pasta(_dados()).save(caminho)
    return load_workbook(caminho)


def test_tem_as_cinco_abas_na_ordem(pasta):
    assert pasta.sheetnames == [
        planilha.ABA_RESUMO, planilha.ABA_LOJAS, planilha.ABA_VENDEDORES,
        planilha.ABA_ALERTAS, planilha.ABA_PREVISAO,
    ]  # fmt: skip


@pytest.mark.parametrize(
    "aba",
    [planilha.ABA_LOJAS, planilha.ABA_VENDEDORES, planilha.ABA_ALERTAS, planilha.ABA_PREVISAO],
)
def test_abas_de_tabela_tem_cabecalho_fixo_e_filtro(pasta, aba):
    ws = pasta[aba]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref and ws.auto_filter.ref.startswith("A1:")


def test_formatos_brasileiros_na_aba_lojas(pasta):
    ws = pasta[planilha.ABA_LOJAS]
    titulos = [c.value for c in ws[1]]
    linha = {t: ws.cell(2, i + 1) for i, t in enumerate(titulos)}
    assert linha["Venda do dia"].number_format == planilha.FORMATO_REAL
    assert linha["Cobertura do mês"].number_format == planilha.FORMATO_PERCENTUAL
    assert linha["Venda do dia"].value == 1500.0
    assert planilha.FORMATO_REAL.startswith('"R$"')
    assert planilha.FORMATO_DATA == "dd/mm/yyyy"


def test_datas_de_alertas_e_previsao_usam_formato_de_data(pasta):
    alerta = pasta[planilha.ABA_ALERTAS].cell(2, 3)
    assert alerta.value.date() == date(2025, 3, 5) and alerta.number_format == "dd/mm/yyyy"
    mes = pasta[planilha.ABA_PREVISAO].cell(2, 1)
    assert mes.number_format == "mm/yyyy"


def test_formatacao_condicional_verde_e_vermelho_nas_coberturas(pasta):
    regras = pasta[planilha.ABA_LOJAS].conditional_formatting
    intervalos = {str(r.sqref): [x.operator for x in r.rules] for r in regras}
    assert intervalos["F2:F3"] == ["greaterThanOrEqual", "lessThan"]  # cobertura do dia
    assert "J2:J3" in intervalos and "N2:N3" in intervalos
    assert "L2:L3" in intervalos  # evolução: verde a partir de zero
    resumo = pasta[planilha.ABA_RESUMO].conditional_formatting
    assert len(list(resumo)) >= 3


def test_alertas_ordenados_como_vieram_e_sem_formulas(pasta):
    ws = pasta[planilha.ABA_ALERTAS]
    assert [ws.cell(i, 1).value for i in (2, 3)] == ["Crítica", "Baixa"]
    for aba in pasta.worksheets:
        for linha in aba.iter_rows():
            for celula in linha:
                assert not (isinstance(celula.value, str) and celula.value.startswith("="))


def test_resumo_traz_indicadores_calculados_dos_dados(pasta):
    ws = pasta[planilha.ABA_RESUMO]
    valores = {ws.cell(i, 1).value: ws.cell(i, 2).value for i in range(4, 20)}
    assert valores["Venda do dia"] == 1900.0 and valores["Meta do dia"] == 1800.0
    assert valores["Cobertura do mês"] == pytest.approx(18000 / 36000)
    assert valores["Acumulado do mês"] == 18000.0
    assert valores["Evolução vs ano anterior"] == pytest.approx(18000 / 17000 - 1)
    assert valores["Projeção de fechamento"] == 37000.0
    assert valores["Cobertura projetada"] == pytest.approx(37000 / 36000)
    assert valores["Ticket médio"] == 300.0 and valores["Cupons do mês"] == 60
    assert "15/03/2025" in ws["A2"].value and "03/2025" in ws["A1"].value


def test_ha_graficos_nativos_no_resumo_e_na_previsao(pasta):
    assert len(pasta[planilha.ABA_RESUMO]._charts) == 1
    assert len(pasta[planilha.ABA_PREVISAO]._charts) == 1
    assert len(pasta[planilha.ABA_LOJAS]._charts) == 0


def test_sem_projecao_nem_previsao_as_celulas_ficam_vazias(tmp_path):
    caminho = tmp_path / "g.xlsx"
    montar_pasta(_dados(com_previsao=False, com_projecao=False)).save(caminho)
    wb = load_workbook(caminho)
    ws = wb[planilha.ABA_RESUMO]
    valores = {ws.cell(i, 1).value: ws.cell(i, 2).value for i in range(4, 20)}
    assert valores["Projeção de fechamento"] is None and valores["Cobertura projetada"] is None
    assert wb[planilha.ABA_PREVISAO].max_row == 1
    assert len(wb[planilha.ABA_PREVISAO]._charts) == 0


def test_venda_depois_do_dia_de_posicao_fica_vazia(pasta):
    ws = pasta[planilha.ABA_RESUMO]
    inicio = next(i for i in range(1, 60) if ws.cell(i, 1).value == "Dia")
    assert ws.cell(inicio + 15, 2).value == 1500.0
    assert ws.cell(inicio + 16, 2).value is None


def test_usa_pasta_existente_e_remove_as_abas_antigas(tmp_path):
    from openpyxl import Workbook

    modelo = Workbook()
    modelo.active.title = "Lixo"
    modelo.create_sheet("Resumo da rede").append(["antigo"])
    montar_pasta(_dados(), modelo)
    assert modelo.sheetnames[0] == planilha.ABA_RESUMO and "Lixo" not in modelo.sheetnames
    assert modelo[planilha.ABA_RESUMO]["A1"].value != "antigo"
    assert len(modelo.sheetnames) == 5
