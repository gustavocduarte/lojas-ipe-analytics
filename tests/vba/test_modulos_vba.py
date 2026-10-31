"""Teste estático dos módulos VBA (o CI não tem Excel): estrutura, segurança e consistência com a
planilha gerada. O roteiro de teste manual está em docs/macro_vba.md."""

import re
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src import config
from src.relatorios import planilha
from src.relatorios.planilha import montar_pasta
from tests.relatorios.test_planilha import _dados

MODULOS = {
    "ModuloLoja": {"GerarAbaDaLoja"},
    "ModuloPDF": {"ExportarResumoPDF"},
    "ModuloAtalhos": {"RegistrarAtalhos", "RemoverAtalhos", "Auto_Open", "Auto_Close"},
}
PROIBIDOS = [
    r"\bShell\b", r"\bKill\b", r"\bCreateObject\b", r"\bADODB\b", r"Provider\s*=", r"\bhttps?:",
    r"\bWorkbooks\.Open\b", r"\bDeclare\b", r"\bSendKeys\b", r"\bURLDownload",
]  # fmt: skip
ABAS_CONHECIDAS = {
    planilha.ABA_RESUMO, planilha.ABA_LOJAS, planilha.ABA_VENDEDORES, planilha.ABA_ALERTAS,
    planilha.ABA_PREVISAO,
}  # fmt: skip


def _ler(nome: str) -> str:
    return (config.DIR_VBA / f"{nome}.bas").read_bytes().decode("ascii")


def _codigo(nome: str) -> list[str]:
    """Linhas sem comentário (o apóstrofo dentro de texto entre aspas não ocorre nestes módulos)."""
    return [re.sub(r"'.*$", "", linha).rstrip() for linha in _ler(nome).splitlines()]


@pytest.mark.parametrize("modulo", MODULOS)
def test_arquivo_pronto_para_importar_no_editor_do_vba(modulo):
    texto = _ler(modulo)  # ASCII puro: importa igual em qualquer página de código do Windows
    linhas = texto.split("\r\n")
    assert linhas[0] == f'Attribute VB_Name = "{modulo}"'
    assert "Option Explicit" in linhas[:3]
    assert "\n" not in texto.replace("\r\n", "")  # só quebras CRLF


@pytest.mark.parametrize("modulo,esperadas", MODULOS.items())
def test_procedimentos_publicos_esperados(modulo, esperadas):
    publicas = set(re.findall(r"^Public Sub (\w+)\(", _ler(modulo), flags=re.M))
    assert publicas == esperadas


@pytest.mark.parametrize("modulo", MODULOS)
def test_blocos_balanceados(modulo):
    codigo = "\n".join(_codigo(modulo))
    aberturas = len(re.findall(r"^\s*(?:Public |Private )?(?:Sub|Function)\s", codigo, flags=re.M))
    assert aberturas == len(re.findall(r"^\s*End (?:Sub|Function)\b", codigo, flags=re.M))
    assert len(re.findall(r"\bThen\s*$", codigo, flags=re.M)) == len(
        re.findall(r"^\s*End If\b", codigo, flags=re.M)
    )
    assert len(re.findall(r"^\s*For\b", codigo, flags=re.M)) == len(
        re.findall(r"^\s*Next\b", codigo, flags=re.M)
    )
    assert len(re.findall(r"^\s*With\b", codigo, flags=re.M)) == len(
        re.findall(r"^\s*End With\b", codigo, flags=re.M)
    )


@pytest.mark.parametrize("modulo", MODULOS)
def test_sem_caminho_absoluto_rede_banco_ou_comando_do_sistema(modulo):
    codigo = "\n".join(_codigo(modulo))
    assert not re.search(r"[A-Za-z]:\\|\\\\\w|/Users/|/home/", codigo)  # caminhos fixos
    for padrao in PROIBIDOS:
        assert not re.search(padrao, codigo, flags=re.I), padrao


def test_abas_citadas_nas_macros_existem_na_planilha_gerada():
    citadas = set()
    for modulo in MODULOS:
        citadas |= set(re.findall(r'Const ABA_\w+ As String = "([^"]+)"', _ler(modulo)))
    assert citadas and citadas <= ABAS_CONHECIDAS


def test_titulos_de_coluna_citados_nas_macros_existem_nas_abas():
    texto = _ler("ModuloLoja")
    titulos = dict(re.findall(r'Const (TITULO_\w+) As String = "([^"]+)"', texto))
    de = lambda colunas: {t for t, _ in colunas.values()}  # noqa: E731
    assert titulos["TITULO_CODIGO_LOJA"] in de(planilha.COLUNAS_LOJAS)
    assert titulos["TITULO_CODIGO_LOJA"] in de(planilha.COLUNAS_VENDEDORES)
    assert titulos["TITULO_NOME_LOJA"] in de(planilha.COLUNAS_LOJAS)
    assert titulos["TITULO_LOJA_NOS_ALERTAS"] in de(planilha.COLUNAS_ALERTAS)


def test_atalhos_apontam_para_macros_que_existem():
    texto = _ler("ModuloAtalhos")
    alvos = re.findall(r'Application\.OnKey "[^"]+", "(\w+)"', texto)
    publicas = {nome for esperadas in MODULOS.values() for nome in esperadas}
    assert alvos and set(alvos) <= publicas
    assert len(re.findall(r'Application\.OnKey "[^"]+"\s*$', texto, flags=re.M)) == len(alvos)


MODELO = Path(config.ARQ_MODELO_XLSM)


@pytest.mark.skipif(not MODELO.exists(), reason="vba/modelo_fechamento.xlsm ainda não foi criado")
def test_xlsm_gerado_a_partir_do_modelo_mantem_o_projeto_vba(tmp_path):
    destino = tmp_path / "fechamento.xlsm"
    montar_pasta(_dados(), load_workbook(MODELO, keep_vba=True)).save(destino)
    with zipfile.ZipFile(destino) as pacote:
        nomes = pacote.namelist()
        assert "xl/vbaProject.bin" in nomes
        assert "macroEnabled" in pacote.read("[Content_Types].xml").decode("utf-8")
    assert load_workbook(destino).sheetnames[0] == planilha.ABA_RESUMO


@pytest.mark.skipif(not MODELO.exists(), reason="vba/modelo_fechamento.xlsm ainda não foi criado")
def test_modelo_traz_os_tres_modulos_importados():
    with zipfile.ZipFile(MODELO) as pacote:
        projeto = pacote.read("xl/vbaProject.bin")
    for modulo in MODULOS:
        assert modulo.encode("utf-16-le") in projeto  # nome do fluxo no arquivo OLE
