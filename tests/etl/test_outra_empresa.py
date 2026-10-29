"""Prova de generalidade: a planilha Excel de "outra empresa" (nomes de coluna diferentes,
cabeçalho na linha 3, linha de total e aba extra) carrega com o mapeamento próprio e dá o
mesmo resultado que o CSV do ERP para o mesmo mês."""

import dataclasses

import pytest
from openpyxl import load_workbook

from src import banco, config
from src.etl import pipeline
from src.etl.leitura import arquivos_da_entidade, ler_registros
from src.etl.mapeamento import carregar_fontes
from src.gerador import outra_empresa
from tests.conftest import banco_temporario
from tests.gerador.utilitarios import ler_csv

ANO, MES = config.MES_EXEMPLO_OUTRA_EMPRESA
ARQUIVO = config.MODELO_ARQ_OUTRA_EMPRESA.format(ano=ANO, mes=MES)
ENTIDADES_DE_CADASTRO = [
    "lojas",
    "produtos",
    "formas_pagamento",
    "campanhas",
    "vendedores",
    "clientes",
]


@pytest.fixture(scope="module")
def planilha(dados_pequenos):
    params, _ = dados_pequenos
    return params.diretorio_saida / config.PASTA_OUTRA_EMPRESA / ARQUIVO


@pytest.fixture(scope="module")
def fontes(dados_pequenos):
    params, _ = dados_pequenos
    todas = carregar_fontes()
    outra = dataclasses.replace(
        todas["outra_empresa"], pasta=params.diretorio_saida / config.PASTA_OUTRA_EMPRESA
    )
    erp = dataclasses.replace(todas["lojas_ipe_erp"], pasta=params.diretorio_saida)
    return erp, outra


def ler_colunas_do_erp():
    return carregar_fontes()["lojas_ipe_erp"].entidades["vendas"].colunas.values()


class TestPlanilhaGerada:
    def test_layout_da_outra_empresa(self, planilha):
        pasta = load_workbook(planilha, read_only=True)
        assert pasta.sheetnames == [outra_empresa.ABA_EXTRA, outra_empresa.ABA_VENDAS]
        linhas = list(pasta[outra_empresa.ABA_VENDAS].iter_rows(values_only=True))
        assert "Vendas de março/2024" in linhas[0][0]
        assert all(c is None for c in linhas[1])
        assert linhas[2][0:2] == ("Nº Cupom", "Item")
        assert "Dt. Emissão" in linhas[2] and "Vlr. Líquido" in linhas[2]
        assert linhas[-1][0] == "TOTAL"
        assert all(c is None for c in linhas[-2])
        pasta.close()

    def test_nao_aparece_no_padrao_dos_csv_do_erp(self, dados_pequenos):
        params, _ = dados_pequenos
        assert not list(params.diretorio_saida.glob("vendas_*.xlsx"))
        assert [p.name for p in (params.diretorio_saida / "outra_empresa").iterdir()] == [ARQUIVO]


class TestMapeamento:
    def test_fonte_registrada_mas_fora_da_carga_padrao(self):
        fonte = carregar_fontes()["outra_empresa"]
        assert fonte.carga_padrao is False
        assert set(fonte.entidades) == {"vendas"}
        mapa = fonte.entidades["vendas"]
        assert (mapa.formato, mapa.aba, mapa.linha_cabecalho) == ("xlsx", "Vendas", 3)
        assert set(mapa.colunas.values()) == set(ler_colunas_do_erp())

    def test_leitor_encontra_os_mesmos_registros_do_csv_do_erp(self, fontes, dados_pequenos):
        params, _ = dados_pequenos
        _, outra = fontes
        mapa = outra.entidades["vendas"]
        contagem = {}
        (arquivo,) = arquivos_da_entidade(outra, mapa)
        registros = list(ler_registros(mapa, arquivo, contagem))
        csv = ler_csv(params.diretorio_saida / f"vendas_{ANO}-{MES:02d}.csv")
        assert contagem == {"lidas": len(csv), "ignoradas": 2}  # linha em branco + total
        assert len(registros) == len(csv)
        colunas = list(mapa.colunas.values())
        primeira = dict(zip(colunas, registros[0], strict=True))
        esperada = csv.iloc[0]
        assert primeira["id_venda_erp"] == esperada["id_venda_erp"]
        assert primeira["numero_cupom"] == esperada["numero_cupom"]
        assert primeira["data_venda"].endswith("/03/2024")
        assert primeira["hora_venda"].count(":") == 2  # horas nativas do Excel viram HH:MM:SS


# A hora é comparada até o minuto: parte das horas do ERP chega no formato HH:MM ou HHhMM, que
# já perde os segundos na origem (a limpeza só pode devolver HH:MM:00). A planilha tem a hora exata.
CONSULTA_VENDAS_DO_MES = """
SELECT f.id_venda_erp, f.numero_cupom, f.numero_item,
       date_trunc('minute', f.data_hora_venda) AS data_hora_ao_minuto, l.cod_loja, f.caixa,
       v.cod_vendedor, c.cod_cliente, fp.cod_forma_pagamento, f.numero_parcelas, cp.cod_campanha,
       p.sku, f.quantidade, f.preco_unitario, f.desconto, f.receita_liquida, f.custo, f.status,
       f.motivo_cancelamento_devolucao
FROM dw.fato_vendas f
JOIN dw.dim_loja l USING (sk_loja)
JOIN dw.dim_vendedor v USING (sk_vendedor)
LEFT JOIN dw.dim_cliente c USING (sk_cliente)
JOIN dw.dim_forma_pagamento fp USING (sk_forma_pagamento)
LEFT JOIN dw.dim_campanha cp USING (sk_campanha)
JOIN dw.dim_produto p USING (sk_produto)
WHERE f.sk_tempo BETWEEN %s AND %s
ORDER BY f.id_venda_erp
"""


@pytest.mark.banco
def test_planilha_carrega_igual_ao_csv_do_erp(banco_com_ddl, fontes, tmp_path):
    erp, outra = fontes
    # Caminho 1: ERP completo (o CSV do mês chega sujo e é corrigido na limpeza).
    pipeline.carregar(banco_com_ddl, erp, verificar_memoria=False, pasta_relatorios=tmp_path)
    esperado = banco_com_ddl.execute(
        CONSULTA_VENDAS_DO_MES, (ANO * 10000 + MES * 100 + 1, ANO * 10000 + MES * 100 + 31)
    ).fetchall()
    assert esperado

    # Caminho 2: só os cadastros do ERP + a planilha da outra empresa, num banco novo.
    with banco_temporario() as outro:
        banco.aplicar_ddl(outro)
        banco.aplicar_procedures(outro)
        pipeline.carregar(
            outro, erp, entidades=ENTIDADES_DE_CADASTRO, verificar_memoria=False,
            pasta_relatorios=tmp_path,
        )  # fmt: skip
        resultado = pipeline.carregar(
            outro, outra, verificar_memoria=False, pasta_relatorios=tmp_path
        )
        assert resultado.aprovada
        assert resultado.relatorio["resumo"]["linhas_rejeitadas"] == 0
        assert resultado.relatorio["arquivos"][0]["linhas_ignoradas"] == 2
        obtido = outro.execute(
            CONSULTA_VENDAS_DO_MES, (ANO * 10000 + MES * 100 + 1, ANO * 10000 + MES * 100 + 31)
        ).fetchall()
        total_no_dw = outro.execute("SELECT count(*) FROM dw.fato_vendas").fetchone()[0]

    assert total_no_dw == len(esperado)  # só o mês da planilha
    assert len(obtido) == len(esperado)
    divergencias = [
        (a[0], [(i, x, y) for i, (x, y) in enumerate(zip(a, b, strict=True)) if x != y])
        for a, b in zip(obtido, esperado, strict=True)
        if a != b
    ]
    assert not divergencias, divergencias[:5]
